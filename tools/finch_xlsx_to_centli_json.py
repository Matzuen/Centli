# -*- coding: utf-8 -*-
"""Convert 「我的账本 / Finch Import」xlsx → Centli centli-export JSON."""
from __future__ import annotations

import argparse
import json
import uuid
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

LEDGER_NAME = "我的账本"
BASE_CCY = "CNY"

ACCOUNT_TYPE = {
    "微信": "ewallet",
    "支付宝": "ewallet",
    "农业银行": "bank",
    "中国银行": "bank",
    "花呗": "credit",
    "白条": "credit",
    "现金": "cash",
}

# 用户截图提供的完整分类树（先写入 JSON，再按流水补齐缺失项）
EXPENSE_TREE: list[tuple[str, list[str]]] = [
    ("餐饮", ["早餐", "午餐", "晚餐", "夜宵", "饮料", "材料", "零食"]),
    ("交通", ["共享", "公交", "地铁", "打车", "能源", "火车", "机票", "停车", "养车", "物流"]),
    ("购物", ["日用品", "服饰", "美妆", "乐意", "数码", "零配件", "保险"]),
    ("娱乐", ["羽毛球", "电影", "游戏", "订阅", "聚会", "健身"]),
    ("健康", ["医疗", "药品", "体检", "保健", "按摩"]),
    ("住房", ["房租", "水电", "物业", "家具", "装修", "燃气"]),
    ("教育", ["学费", "课程", "书籍", "文具", "考试"]),
    ("通讯", ["话费", "网费"]),
    ("宠物", ["宠物食品", "宠物用品", "宠物护理"]),
    ("社交", ["礼物", "请客", "礼金"]),
    ("旅行", ["住宿", "门票", "旅行杂费"]),
    ("金融费用", ["手续费", "托管费", "汇损"]),
    ("其他支出", ["捐赠", "丢失", "其他"]),
]

INCOME_TREE: list[tuple[str, list[str]]] = [
    ("主业", ["薪资", "奖金"]),
    ("副业", ["短视频", "小说", "回血", "兼职"]),
    ("投资", ["利息", "回报"]),
    ("红包", ["天降", "人为", "人情往来"]),
    ("理赔", []),
]


def new_id() -> str:
    return str(uuid.uuid4())


def yuan_to_minor(v) -> int:
    if v is None or v == "":
        return 0
    return int(round(float(v) * 100))


def parse_happen_at(s: str) -> int:
    """Parse 'YYYY-MM-DD HH:MM' as Asia/Shanghai local time → epoch ms."""
    from calendar import timegm

    dt = datetime.strptime(str(s).strip(), "%Y-%m-%d %H:%M")
    # timegm 把 tuple 当 UTC；此处数字实为东八区墙钟 → 减 8 小时得 UTC。
    return (timegm(dt.timetuple()) - 8 * 3600) * 1000


def account_type_for(name: str) -> str:
    return ACCOUNT_TYPE.get(name, "other")


def build_note(row: dict) -> str | None:
    parts: list[str] = []
    for key in ("description", "merchant", "sub_category"):
        v = (row.get(key) or "").strip()
        if v:
            parts.append(v)
    member = (row.get("member") or "").strip()
    if member:
        parts.append(f"成员:{member}")
    if not parts:
        return None
    return " · ".join(parts)


def build_tags(row: dict) -> str | None:
    tags: list[str] = []
    raw = (row.get("tag") or "").strip()
    if raw:
        for piece in raw.replace("，", ",").split(","):
            t = piece.strip()
            if t and t not in tags:
                tags.append(t)
    sub = (row.get("sub_category") or "").strip()
    if sub in ("退款", "返现") and "退款" not in tags:
        tags.append("退款")
    if sub == "还款" and "还款" not in tags:
        tags.append("还款")
    if not tags:
        return None
    return json.dumps(tags, ensure_ascii=False)


def map_type(finch_type: str) -> str:
    return {
        "expense": "out",
        "income": "in",
        "transfer": "transfer",
    }.get(str(finch_type).strip().lower(), "out")


def convert(xlsx: Path) -> dict:
    wb = load_workbook(xlsx, data_only=True)
    if "Finch Import" not in wb.sheetnames:
        raise SystemExit(f"缺少「Finch Import」工作表，现有: {wb.sheetnames}")
    ws = wb["Finch Import"]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    rows = [
        dict(zip(headers, r))
        for r in ws.iter_rows(min_row=2, values_only=True)
        if any(x is not None and str(x).strip() != "" for x in r)
    ]
    wb.close()

    now = int(datetime.now().timestamp() * 1000)
    ledger_id = new_id()

    # ---- accounts ----
    account_names: set[str] = set()
    for r in rows:
        for k in ("from_account", "to_account"):
            n = (r.get(k) or "").strip()
            if n:
                account_names.add(n)
    accounts = []
    account_ids: dict[str, str] = {}
    for i, name in enumerate(sorted(account_names)):
        aid = new_id()
        account_ids[name] = aid
        accounts.append(
            {
                "id": aid,
                "ledgerId": ledger_id,
                "name": name,
                "currency": BASE_CCY,
                "type": account_type_for(name),
                "icon": None,
                "openingBalanceMinor": 0,
                "sortOrder": i,
                "archived": 0,
                "createdAt": now,
                "updatedAt": now,
            }
        )

    # ---- categories (parent + child per kind) ----
    # key: (kind, parent_name) / (kind, parent_name, child_name)
    parent_ids: dict[tuple[str, str], str] = {}
    child_ids: dict[tuple[str, str, str], str] = {}
    categories: list[dict] = []
    parent_order = 0
    child_order = 0

    def ensure_parent(kind: str, name: str) -> str:
        nonlocal parent_order
        name = (name or "").strip() or ("支出" if kind == "out" else "收入")
        key = (kind, name)
        if key in parent_ids:
            return parent_ids[key]
        pid = new_id()
        parent_ids[key] = pid
        categories.append(
            {
                "id": pid,
                "ledgerId": ledger_id,
                "name": name,
                "parentId": None,
                "kind": kind,
                "icon": None,
                "sortOrder": parent_order,
                "archived": 0,
                "createdAt": now,
            }
        )
        parent_order += 1
        return pid

    def ensure_child(kind: str, parent: str, child: str) -> str | None:
        nonlocal child_order
        child = (child or "").strip()
        if not child:
            return ensure_parent(kind, parent)
        key = (kind, parent or "", child)
        if key in child_ids:
            return child_ids[key]
        pid = ensure_parent(kind, parent)
        cid = new_id()
        child_ids[key] = cid
        categories.append(
            {
                "id": cid,
                "ledgerId": ledger_id,
                "name": child,
                "parentId": pid,
                "kind": kind,
                "icon": None,
                "sortOrder": child_order,
                "archived": 0,
                "createdAt": now,
            }
        )
        child_order += 1
        return cid

    # 先写入截图完整分类树（含尚未出现在流水里的子分类）
    for parent, children in EXPENSE_TREE:
        ensure_parent("out", parent)
        for child in children:
            ensure_child("out", parent, child)
    for parent, children in INCOME_TREE:
        ensure_parent("in", parent)
        for child in children:
            ensure_child("in", parent, child)

    # ---- transactions ----
    # Import every non-zero row (main + sub). Related refund/AA/fee stay separate.
    transactions: list[dict] = []
    skipped_zero = 0
    for r in rows:
        amount = r.get("amount")
        if amount is None:
            amount = r.get("currency_amount")
        minor = yuan_to_minor(amount)
        if minor == 0:
            skipped_zero += 1
            continue

        t = map_type(r.get("type") or "expense")
        from_name = (r.get("from_account") or "").strip()
        to_name = (r.get("to_account") or "").strip()
        if not from_name:
            # income sometimes only has from_account as the receiving wallet
            continue
        if from_name not in account_ids:
            continue

        cat_id = None
        if t != "transfer":
            kind = "out" if t == "out" else "in"
            cat_id = ensure_child(
                kind,
                str(r.get("parent_category") or ""),
                str(r.get("category") or ""),
            )

        ccy = (r.get("currency") or r.get("book_currency") or BASE_CCY).strip() or BASE_CCY
        happen = parse_happen_at(str(r.get("date_time")))
        note = build_note(r)
        tags = build_tags(r)
        include_stat = 1 if str(r.get("include_statistics") or "1").strip() in ("1", "true", "True") else 0
        include_budget = 1 if str(r.get("include_budget") or "1").strip() in ("1", "true", "True") else 0

        txn = {
            "id": new_id(),
            "ledgerId": ledger_id,
            "type": t,
            "amountMinor": minor,
            "currency": ccy,
            "accountId": account_ids[from_name],
            "toAccountId": account_ids.get(to_name) if t == "transfer" and to_name else None,
            "categoryId": cat_id,
            "note": note,
            "tags": tags,
            "happenAt": happen,
            "includeInStat": include_stat,
            "includeInBudget": include_budget,
            "transferGroupId": None,
            "fxRate": None,
            "toAmountMinor": minor if t == "transfer" else None,
            "source": "import",
            "pendingId": None,
            "occurrenceKey": f"finch|{r.get('group_key')}|{r.get('level')}",
            "createdAt": now,
            "updatedAt": now,
            "deletedAt": None,
        }
        if t == "transfer" and to_name and to_name in account_ids:
            txn["transferGroupId"] = new_id()
        transactions.append(txn)

    # Parents first already ensured by ensure_parent before children.
    # Stable sort: null parentId first.
    categories.sort(key=lambda c: (0 if c["parentId"] is None else 1, c["kind"], c["sortOrder"]))

    bundle = {
        "format": "centli-export",
        "export_version": 1,
        "schema_version": 4,
        "exported_at": datetime.now().strftime("%Y-%m-%dT%H:%M:%S.%f") + "Z",
        "app_version": "import-from-finch",
        "ledger": {
            "id": ledger_id,
            "name": LEDGER_NAME,
            "baseCurrency": BASE_CCY,
            "createdAt": now,
            "archived": 0,
        },
        "accounts": accounts,
        "asset_types": [],
        "categories": categories,
        "transactions": transactions,
        "merchant_rules": [],
        "pending_captures": [],
        "fx_rates": [],
        "budgets": [],
        "settings": {
            "onboarding_done": "true",
            "base_currency": BASE_CCY,
            "auto_capture_enabled": "false",
        },
        "_meta": {
            "source": str(xlsx),
            "source_rows": len(rows),
            "imported_transactions": len(transactions),
            "skipped_zero_amount": skipped_zero,
            "accounts": len(accounts),
            "categories": len(categories),
        },
    }
    return bundle


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx", type=Path)
    ap.add_argument("-o", "--output", type=Path, required=True)
    args = ap.parse_args()
    bundle = convert(args.xlsx)
    meta = bundle.pop("_meta")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(bundle, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(json.dumps(meta, ensure_ascii=False, indent=2))
    print(f"wrote {args.output} ({args.output.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
